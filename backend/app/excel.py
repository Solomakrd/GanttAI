"""The Excel boundary: task-name references outside, ID references inside."""

import csv
import io
import re
from collections import deque
from datetime import date, datetime, timedelta
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_EXPANDED_BYTES = 20 * 1024 * 1024
MAX_ROWS = 500
MAX_SHEET_ROWS = 1_048_576  # Excel's physical worksheet limit, including empty rows.
MAX_DAYS = 730
MAX_COLUMNS = 32
MAX_LINKS = 10000
HEADERS = ["задача", "описание", "исполнитель", "длительность", "предшественники"]
OPTIONAL_HEADERS = ["id", "start_date", "end_date"]


class ExcelError(Exception):
    def __init__(self, message, *, sheet=None, row=None, column=None, status=422):
        self.detail = {"message": message, "sheet": sheet, "row": row, "column": column}
        self.status = status
        super().__init__(message)


def fail(message, **location):
    raise ExcelError(message, **location)


def text(value, column, location, required=False, trim=False):
    if value is None:
        value = ""
    if not isinstance(value, str):
        fail("Use a text cell.", column=column, **location)
    if len(value) > 32767 or ILLEGAL_CHARACTERS_RE.search(value):
        fail("Text is too long or contains unsupported control characters.", column=column, **location)
    if trim:
        value = value.strip()
    if required and not value.strip():
        fail("A non-empty value is required.", column=column, **location)
    return value


def serialize_predecessors(names):
    output = io.StringIO()
    csv.writer(output, lineterminator="\r\n").writerow(names)
    return output.getvalue().removesuffix("\r\n")


def parse_date(value, column, location):
    if isinstance(value, datetime) and value.time() == datetime.min.time():
        return value.date()
    if type(value) is date:
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        try:
            return date.fromisoformat(value)
        except ValueError:
            pass
    fail("Use a valid date in YYYY-MM-DD format or an Excel date without a time.", column=column, **location)


def validate_schedule(tasks, locations, project_start=None):
    """Validate the graph iteratively, then preserve or calculate inclusive dates."""
    if len(tasks) > MAX_ROWS:
        fail(f"Use at most {MAX_ROWS} tasks.")
    by_id, names = {}, {}
    for task, location in zip(tasks, locations):
        name = text(task["task"], "задача", location, required=True, trim=True)
        if name in names:
            fail(f"Duplicate task name {name!r}; conflicting rows {names[name]} and {location['row']}.", column="задача", **location)
        names[name] = location["row"]
        if task["id"] in by_id:
            fail(f"Duplicate id {task['id']!r}.", column="id", **location)
        by_id[task["id"]] = task
        if not 1 <= task["duration"] <= MAX_DAYS:
            fail(f"Duration must be a whole number from 1 to {MAX_DAYS}.", column="длительность", **location)

    children = {key: [] for key in by_id}
    pending = {}
    if sum(len(task["predecessors"]) for task in tasks) > MAX_LINKS:
        fail(f"Use at most {MAX_LINKS} predecessor relationships.")
    for task, location in zip(tasks, locations):
        predecessors = task["predecessors"]
        if len(set(predecessors)) != len(predecessors):
            fail(f"Repeated predecessor for {task['task']!r}.", column="предшественники", **location)
        pending[task["id"]] = len(predecessors)
        for predecessor in predecessors:
            if predecessor not in by_id:
                fail(f"Unknown predecessor {predecessor!r} for {task['task']!r}.", column="предшественники", **location)
            if predecessor == task["id"]:
                fail(f"Task {task['task']!r} cannot depend on itself.", column="предшественники", **location)
            children[predecessor].append(task["id"])

    ready = deque(key for key, count in pending.items() if count == 0)
    ordered = []
    while ready:
        key = ready.popleft()
        ordered.append(key)
        for child in children[key]:
            pending[child] -= 1
            if not pending[child]:
                ready.append(child)
    if len(ordered) != len(tasks):
        affected = ", ".join(by_id[key]["task"] for key, count in pending.items() if count)
        fail(f"Dependency cycle; affected tasks: {affected}.", column="предшественники")

    location_by_id = {task["id"]: location for task, location in zip(tasks, locations)}
    for key in ordered:
        task = by_id[key]
        location = location_by_id[key]
        latest = max((by_id[p]["end_date"] for p in task["predecessors"]), default=None)
        if task["start_date"] is None:
            if project_start is None:
                fail("Choose a project start date for tasks without supplied dates.", column="start_date", **location)
            try:
                task["start_date"] = latest + timedelta(days=1) if latest else project_start
                task["end_date"] = task["start_date"] + timedelta(days=task["duration"] - 1)
            except OverflowError:
                fail("Calculated dates exceed the supported calendar. Choose an earlier start or shorter duration.", **location)
        if (task["end_date"] - task["start_date"]).days + 1 != task["duration"]:
            fail("Supplied dates must match the inclusive duration.", column="end_date", **location)
        if latest and task["start_date"] <= latest:
            fail(f"Task {task['task']!r} must start after all predecessors finish; supplied dates are not rescheduled.", column="start_date", **location)
    if tasks and (max(t["end_date"] for t in tasks) - min(t["start_date"] for t in tasks)).days + 1 > MAX_DAYS:
        fail(f"The displayed plan must span at most {MAX_DAYS} calendar days.")


def validate_archive(data):
    if len(data) > MAX_FILE_BYTES:
        fail("Workbook exceeds the 2 MiB file limit.", status=413)
    with ZipFile(io.BytesIO(data)) as archive:
        if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
            fail("Expanded workbook exceeds the 20 MiB / 1000-entry limit.", status=413)


def read_workbook(data, project_start=None):
    try:
        validate_archive(data)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ExcelError:
        raise
    except Exception as error:
        raise ExcelError("Cannot read this workbook. Upload a valid, unencrypted .xlsx file.") from error
    try:
        return read_sheet(workbook, project_start)
    except ExcelError:
        raise
    except Exception as error:
        raise ExcelError("Workbook contents are damaged or unsupported. Save it as .xlsx and retry.") from error
    finally:
        workbook.close()


def read_sheet(workbook, project_start):
    if not workbook.worksheets:
        fail("Workbook must contain a worksheet.")
    sheet = workbook.worksheets[0]
    # Ignore untrusted worksheet dimensions; bound populated rows and columns below.
    sheet.reset_dimensions()
    rows = sheet.iter_rows()
    header_cells = next(rows, ())
    if any(cell.data_type in ("f", "e") for cell in header_cells):
        fail("Formulas and Excel error cells are not accepted in headers; paste values instead.", sheet=sheet.title, row=1)
    if len(header_cells) > MAX_COLUMNS:
        fail(f"Use at most {MAX_COLUMNS} columns.", sheet=sheet.title, row=1)
    headers = [cell.value.strip() if isinstance(cell.value, str) else cell.value for cell in header_cells]
    missing = [name for name in HEADERS if name not in headers]
    if missing:
        fail(f"Missing required headers: {', '.join(missing)}.", sheet=sheet.title, row=1)
    for name in HEADERS + OPTIONAL_HEADERS:
        if headers.count(name) > 1:
            fail(f"Duplicate header {name!r}.", sheet=sheet.title, row=1, column=name)
    tasks, locations, name_index, references = [], [], {}, []
    for row_number, cells in enumerate(rows, start=2):
        location = {"sheet": sheet.title, "row": row_number}
        if row_number > MAX_SHEET_ROWS:
            fail(f"Worksheet exceeds Excel's {MAX_SHEET_ROWS}-row limit.", **location)
        if len(cells) > MAX_COLUMNS:
            fail(f"Use at most {MAX_COLUMNS} columns.", **location)
        for cell in cells:
            if cell.data_type in ("f", "e"):
                fail("Formulas and Excel error cells are not accepted; paste values instead.", column=cell.column_letter, **location)
        values = [cell.value for cell in cells]
        if all(value is None or (isinstance(value, str) and not value.strip()) for value in values):
            continue
        if len(tasks) >= MAX_ROWS:
            fail(f"Use at most {MAX_ROWS} data rows (empty rows are ignored).", **location)
        record = {name: values[index] if index < len(values) else None for index, name in enumerate(headers) if isinstance(name, str)}
        name = text(record.get("задача"), "задача", location, required=True, trim=True)
        if name in name_index:
            fail(f"Duplicate task name {name!r}; conflicting rows {name_index[name]} and {row_number}.", column="задача", **location)
        name_index[name] = row_number
        duration = record.get("длительность")
        if isinstance(duration, bool) or not isinstance(duration, (int, float)) or not 1 <= duration <= MAX_DAYS or int(duration) != duration:
            fail(f"Duration must be a whole number from 1 to {MAX_DAYS}.", column="длительность", **location)
        predecessor_text = text(record.get("предшественники"), "предшественники", location)
        try:
            parsed = list(csv.reader(io.StringIO(predecessor_text), strict=True, skipinitialspace=True)) if predecessor_text.strip() else [[]]
            if len(parsed) != 1 or any(not value.strip() for value in parsed[0]):
                raise ValueError()
            references.append([value.strip() for value in parsed[0]])
        except (csv.Error, ValueError):
            fail('Use comma-separated task names; quote names containing commas or quotes using CSV syntax.', column="предшественники", **location)
        start, end = record.get("start_date"), record.get("end_date")
        if (start is None) != (end is None):
            fail("Supply both start_date and end_date, or leave both blank.", column="start_date/end_date", **location)
        tasks.append({
            "id": text(record.get("id"), "id", location) or None,
            "task": name,
            "description": text(record.get("описание"), "описание", location),
            "assignee": text(record.get("исполнитель"), "исполнитель", location, required=True),
            "duration": int(duration),
            "start_date": parse_date(start, "start_date", location) if start is not None else None,
            "end_date": parse_date(end, "end_date", location) if end is not None else None,
            "predecessors": [],
        })
        locations.append(location)
    used_ids = {task["id"] for task in tasks if task["id"] is not None}
    for task, location in zip(tasks, locations):
        if task["id"] is None:
            key = f"task-{location['row'] - 1}"
            while key in used_ids:
                key += "-new"
            task["id"] = key
            used_ids.add(key)
    ids_by_name = {task["task"]: task["id"] for task in tasks}
    for task, names, location in zip(tasks, references, locations):
        for name in names:
            if name not in ids_by_name:
                fail(f"Unknown predecessor {name!r} for {task['task']!r}; names are case-sensitive.", column="предшественники", **location)
        text(serialize_predecessors(names), "предшественники", location)
        task["predecessors"] = [ids_by_name[name] for name in names]
    validate_schedule(tasks, locations, project_start)
    return {"tasks": tasks}


def write_workbook(tasks):
    locations = [{"sheet": "Plan", "row": i + 2} for i in range(len(tasks))]
    validate_schedule(tasks, locations)
    names = {task["id"]: task["task"].strip() for task in tasks}
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan"
    sheet.append(HEADERS + OPTIONAL_HEADERS)
    for task, location in zip(tasks, locations):
        predecessors = serialize_predecessors([names[key] for key in task["predecessors"]])
        assignee = text(task["assignee"], "исполнитель", location, required=True)
        values = [task["task"].strip(), task["description"], assignee, task["duration"], predecessors, task["id"], task["start_date"].isoformat(), task["end_date"].isoformat()]
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=location["row"], column=index)
            if isinstance(value, str):
                text(value, (HEADERS + OPTIONAL_HEADERS)[index - 1], location)
                cell.value = value
                cell.data_type = "s"  # Even '=...', '+...', etc. are literal text.
            else:
                cell.value = value
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    data = output.getvalue()
    validate_archive(data)
    return data
