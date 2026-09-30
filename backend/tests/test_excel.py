import csv
import io
from datetime import date, datetime
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import excel
from app.main import app

client = TestClient(app)


def workbook_bytes(rows, headers=None, second_sheet=False):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tasks"
    sheet.append(headers if headers is not None else excel.HEADERS)
    for row in rows:
        sheet.append(row)
    if second_sheet:
        workbook.create_sheet("Ignored").append(["not a valid task sheet"])
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def upload(data, start="2026-10-02", filename="tasks.xlsx"):
    return client.post("/api/plan/import", files={"file": (filename, data)}, data={"start_date": start} if start else {})


def assert_error(response, message, row=None, column=None):
    assert response.status_code in (413, 422), response.text
    detail = response.json()["detail"]
    assert message.lower() in detail["message"].lower()
    if row is not None:
        assert detail["row"] == row
        assert detail["sheet"] == "Tasks"
    if column:
        assert detail["column"] == column


def test_trailing_formatted_empty_rows_do_not_count_as_tasks():
    rows = [[f"Task {i}", "", "A", 1, ""] for i in range(5)]
    workbook = load_workbook(io.BytesIO(workbook_bytes(rows)))
    sheet = workbook.active
    sheet.cell(1000, 1).number_format = "@"
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    result = upload(stream.getvalue())
    assert result.status_code == 200, result.text
    assert [t["task"] for t in result.json()["tasks"]] == [r[0] for r in rows]


def test_malformed_xml_row_index_cannot_expand_unbounded_empty_rows(monkeypatch):
    monkeypatch.setattr(excel, "MAX_SHEET_ROWS", 1000)
    data = workbook_bytes([["Root", "", "A", 1, ""]])
    output = io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'<row r="2">', b'<row r="1000000000000">')
            target.writestr(info, content)
    assert_error(upload(output.getvalue()), "Excel's 1000-row limit", row=1001)


def test_task_limit_ignores_gaps_but_reports_original_overflow_row():
    rows = []
    for i in range(excel.MAX_ROWS):
        rows.extend([[" ", "\t", None, None, None], [f"Task {i}", "", "A", 1, ""]])
    result = upload(workbook_bytes(rows))
    assert result.status_code == 200, result.text
    assert len(result.json()["tasks"]) == excel.MAX_ROWS
    rows.extend([[], ["Extra", "", "A", 1, ""]])
    assert_error(upload(workbook_bytes(rows)), "500 data rows", row=1003)


@pytest.mark.parametrize("row, message, column", [
    (["Invalid", "", "A", 0, ""], "whole number", "длительность"),
    ([None, '=IF(1=1,"","")'], "Formulas", "B"),
    ([None, None, None, 0], "non-empty", "задача"),
])
def test_invalid_populated_rows_after_empty_gap_still_validate(row, message, column):
    rows = [["Root", "", "A", 1, ""]] + [[] for _ in range(600)] + [row]
    assert_error(upload(workbook_bytes(rows)), message, row=603, column=column)


def test_forward_references_inclusive_weekend_and_latest_predecessor():
    result = upload(workbook_bytes([
        ["Successor", "", "B", 2, "Root, Other"],
        ["Root", "Friday through Sunday", "A", 3, ""],
        ["Other", "", "C", 1, ""],
    ], second_sheet=True))
    assert result.status_code == 200, result.text
    successor, root, other = result.json()["tasks"]
    assert (root["start_date"], root["end_date"]) == ("2026-10-02", "2026-10-04")
    assert (successor["start_date"], successor["end_date"]) == ("2026-10-05", "2026-10-06")
    assert successor["predecessors"] == [root["id"], other["id"]]


def test_scheduling_requires_a_chosen_date_and_reports_bad_date():
    data = workbook_bytes([["A", "", "A", 1, ""]])
    assert_error(upload(data, None), "Choose a project start date", row=2)
    result = upload(data, "2026-02-30")
    assert result.status_code == 422
    assert result.json()["detail"][0]["loc"] == ["body", "start_date"]


@pytest.mark.parametrize("value", [0, -1, 1.5, "three", "2", True, None, 731])
def test_invalid_durations(value):
    assert_error(upload(workbook_bytes([["A", "", "Person", value, ""]])), "whole number", row=2, column="длительность")


@pytest.mark.parametrize("rows, message", [
    ([["A", "", "A", 1, "Missing"]], "Unknown predecessor 'Missing'"),
    ([["A", "", "A", 1, "a"]], "case-sensitive"),
    ([["A", "", "A", 1, "A"]], "itself"),
    ([["A", "", "A", 1, "B"], ["B", "", "B", 1, "A"]], "cycle; affected tasks: A, B"),
    ([["A", "", "A", 1, "B, B"], ["B", "", "B", 1, ""]], "Repeated predecessor"),
    ([["A", "", "A", 1, '"unclosed']], "CSV syntax"),
    ([["A", "", "A", 1, "B,"]], "CSV syntax"),
])
def test_graph_and_csv_errors(rows, message):
    assert_error(upload(workbook_bytes(rows)), message)


def test_duplicate_names_report_both_rows_and_case_sensitive_names_are_distinct():
    assert_error(upload(workbook_bytes([[" A ", "", "A", 1, ""], ["A", "", "B", 1, ""]])), "conflicting rows 2 and 3", row=3)
    result = upload(workbook_bytes([[" A ", "", "A", 1, ""], ["a", "", "B", 1, " A "]]))
    assert result.status_code == 200
    assert result.json()["tasks"][0]["task"] == "A"


@pytest.mark.parametrize("data, filename, message", [
    (b"not a zip", "broken.xlsx", "Cannot read"),
    (b"not excel", "tasks.csv", ".xlsx"),
    (workbook_bytes([], headers=["task"]), "tasks.xlsx", "Missing required headers"),
    (workbook_bytes([], headers=excel.HEADERS + ["задача"]), "tasks.xlsx", "Duplicate header"),
    (workbook_bytes([["A", "=1+1", "A", 1, ""]]), "tasks.xlsx", "Formulas"),
    (workbook_bytes([["A", "#REF!", "A", 1, ""]]), "tasks.xlsx", "error cells"),
    (workbook_bytes([], headers=excel.HEADERS + ["=1+1"]), "tasks.xlsx", "headers"),
    (workbook_bytes([[None, "", "A", 1, ""]]), "tasks.xlsx", "non-empty"),
    (workbook_bytes([["A", "", " ", 1, ""]]), "tasks.xlsx", "non-empty"),
])
def test_invalid_workbooks(data, filename, message):
    assert_error(upload(data, filename=filename), message)


@pytest.mark.parametrize("start, end, duration, message", [
    ("2026-02-30", "2026-03-01", 1, "valid date"),
    ("2026-10-02", "2026-10-01", 1, "inclusive duration"),
    ("2026-10-02", "2026-10-04", 2, "inclusive duration"),
    ("2026-10-02", None, 1, "both start_date and end_date"),
    (datetime(2026, 10, 2, 12), datetime(2026, 10, 2), 1, "without a time"),
])
def test_invalid_supplied_dates(start, end, duration, message):
    assert_error(upload(workbook_bytes([["A", "", "A", duration, "", start, end]], excel.HEADERS + ["start_date", "end_date"])), message, row=2)


def test_preserves_dates_and_gaps_and_rejects_overlapping_dependency():
    headers = excel.HEADERS + ["start_date", "end_date"]
    rows = [["A", "", "A", 1, "", date(2026, 10, 2), date(2026, 10, 2)], ["B", "", "B", 1, "A", "2026-10-10", "2026-10-10"]]
    result = upload(workbook_bytes(rows, headers), start=None)
    assert result.status_code == 200, result.text
    assert result.json()["tasks"][1]["start_date"] == "2026-10-10"
    rows[1][-2:] = ["2026-10-02", "2026-10-02"]
    assert_error(upload(workbook_bytes(rows, headers)), "not rescheduled", row=3)


def test_optional_ids_preserved_generated_without_collisions_and_duplicates_rejected():
    rows = [["A", "", "A", 1, "", None], ["B", "", "B", 1, "A", "task-1"]]
    result = upload(workbook_bytes(rows, excel.HEADERS + ["id"]))
    assert result.status_code == 200
    first, second = result.json()["tasks"]
    assert first["id"] != second["id"] == "task-1"
    assert second["predecessors"] == [first["id"]]
    rows[0][-1] = "task-1"
    assert_error(upload(workbook_bytes(rows, excel.HEADERS + ["id"])), "Duplicate id", row=3)


def test_real_file_expanded_size_row_column_and_span_limits():
    assert_error(upload(b"x" * (excel.MAX_FILE_BYTES + 1)), "2 MiB")
    assert_error(upload(b"x" * (excel.MAX_FILE_BYTES + 65537)), "multipart overhead")
    stream = io.BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"x" * (excel.MAX_EXPANDED_BYTES + 1))
    assert_error(upload(stream.getvalue()), "Expanded workbook")
    rows = [[f"Task {i}", "", "A", 1, ""] for i in range(excel.MAX_ROWS)]
    assert upload(workbook_bytes(rows)).status_code == 200
    assert_error(upload(workbook_bytes(rows + [["Extra", "", "A", 1, ""]])), "500 data rows")
    assert_error(upload(workbook_bytes([], excel.HEADERS + [f"extra{i}" for i in range(28)])), "32 columns")
    assert upload(workbook_bytes([["A", "", "A", 730, ""]])).status_code == 200
    assert_error(upload(workbook_bytes([["A", "", "A", 730, ""], ["B", "", "B", 1, "A"]])), "span at most 730")
    assert_error(upload(workbook_bytes([["A", "", "A", 2, ""]]), "9999-12-31"), "exceed the supported calendar")


def test_deep_cycle_is_iterative_not_recursive():
    rows = [[f"T{i}", "", "A", 1, f"T{(i + 1) % 500}"] for i in range(500)]
    assert_error(upload(workbook_bytes(rows)), "cycle")


@pytest.mark.parametrize("separator", ["\n", "\r", "\r\n"])
def test_blank_rows_mixed_dates_and_csv_newlines(separator):
    headers = excel.HEADERS + ["id", "start_date", "end_date"]
    data = workbook_bytes([
        [f'Root{separator}with newline', '', 'A', 1, '', 'root', '2026-10-02', '2026-10-02'],
        [],
        ['Next', '', 'B', 1, f'"Root{separator}with newline"'],
    ], headers)
    result = upload(data)
    assert result.status_code == 200, result.text
    assert result.json()["tasks"][1]["start_date"] == '2026-10-03'
    round_trip(result.json())


def test_link_limit_and_archive_entry_limit(monkeypatch):
    monkeypatch.setattr(excel, 'MAX_LINKS', 1)
    assert_error(upload(workbook_bytes([["A", "", "A", 1, ""], ["B", "", "B", 1, "A"], ["C", "", "C", 1, "B"]])), "predecessor relationships")
    stream = io.BytesIO()
    with ZipFile(stream, 'w', ZIP_DEFLATED) as archive:
        for i in range(1001):
            archive.writestr(str(i), '')
    assert_error(upload(stream.getvalue()), "1000-entry")


def round_trip(plan):
    response = client.post("/api/plan/export", json=plan)
    assert response.status_code == 200, response.text
    assert "spreadsheetml" in response.headers["content-type"]
    assert 'filename="gantt-plan.xlsx"' in response.headers["content-disposition"]
    imported = upload(response.content, "2030-01-01")
    assert imported.status_code == 200, imported.text
    assert imported.json() == plan
    return response.content


def test_seed_and_empty_plan_round_trips():
    round_trip(client.get("/api/plan").json())
    round_trip({"tasks": []})


def test_imported_csv_names_and_literal_formula_text_round_trip():
    csv_names = io.StringIO()
    csv.writer(csv_names, lineterminator="").writerow(['Design, review', 'Say "go"'])
    imported = upload(workbook_bytes([
        ["Ship", "", "A", 1, csv_names.getvalue()],
        ["Design, review", "", "A", 1, ""],
        ['Say "go"', "", "A", 1, ""],
    ])).json()
    imported["tasks"][0]["description"] = '=HYPERLINK("https://example.test")'
    imported["tasks"][0]["assignee"] = "+literal"
    exported = round_trip(imported)
    workbook = load_workbook(io.BytesIO(exported))
    sheet = workbook.active
    assert sheet.cell(2, 2).data_type == "s"
    assert next(csv.reader([sheet.cell(2, 5).value])) == ['Design, review', 'Say "go"']
    workbook.close()


@pytest.mark.parametrize("quote_count", [16383, 17000])
def test_import_rejects_predecessors_that_exceed_cell_limit_after_csv_escaping(quote_count):
    name = "A" + '"' * quote_count
    result = upload(workbook_bytes([
        [name, "", "A", 1, ""],
        ["Successor", "", "B", 1, name],
    ]))
    assert result.status_code == 422
    assert_error(result, "too long", row=3, column="предшественники")
    assert "tasks" not in result.json()


def test_predecessor_csv_at_cell_limit_imports_and_round_trips():
    name = "A" + '"' * 16382  # CSV escaping produces exactly 32,767 characters.
    result = upload(workbook_bytes([
        [name, "", "A", 1, ""],
        ["Successor", "", "B", 1, name],
    ]))
    assert result.status_code == 200, result.text
    round_trip(result.json())


def test_imports_never_mutate_seed_or_other_clients_on_success_or_error():
    seed = client.get("/api/plan").json()
    assert upload(workbook_bytes([["Private", "", "A", 1, ""]])).status_code == 200
    assert_error(upload(workbook_bytes([["Bad", "", "A", 0, ""]])), "Duration")
    assert client.get("/api/plan").json() == seed
    assert TestClient(app).get("/api/plan").json() == seed


@pytest.mark.parametrize("change, message", [
    ({"predecessors": ["missing"]}, "Unknown predecessor"),
    ({"duration": 2}, "inclusive duration"),
    ({"description": "bad\x01text"}, "control characters"),
    ({"description": "x" * 32768}, "too long"),
    ({"assignee": " "}, "non-empty"),
])
def test_invalid_export_returns_actionable_error(change, message):
    plan = upload(workbook_bytes([["A", "", "A", 1, ""]])).json()
    plan["tasks"][0].update(change)
    assert_error(client.post("/api/plan/export", json=plan), message)


def test_multipart_missing_file_and_invalid_export_schema():
    assert client.post("/api/plan/import", data={"start_date": "2026-10-02"}).status_code == 422
    assert client.post("/api/plan/export", json={"tasks": [{"duration": 0}]}).status_code == 422
